from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .common import (
    active_profile,
    choose_output_path,
    configure_utf8_stdio,
    read_json,
    signature_style_spec,
    unique_output_path,
)
from .inspect_docx import (
    analyze_tables, classify_paragraphs, inspect_document, is_spacing_paragraph,
    normalize_legacy_roles, signature_left_indents, split_signature_text, title_block_layout,
)
from .privacy_scrub import scrub_docx
from .validate_docx import validate_document


ROLE_NAMES = {
    "main_title": "ODF Main Title",
    "heading1": "ODF Heading 1",
    "heading2": "ODF Heading 2",
    "body": "ODF Body",
    "reference_note": "ODF Reference Note",
    "description": "ODF Description",
    "signature": "ODF Signature",
}
ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
H1_RE = re.compile(r"^[一二三四五六七八九十百]+、")
H2_RE = re.compile(r"^（[一二三四五六七八九十百]+）")
NOTE_RE = re.compile(r"^(?:（.*）|\(.*\))$", re.DOTALL)


def installed_font_names() -> set[str]:
    names: set[str] = set()
    if sys.platform != "win32":
        return names
    try:
        import winreg
    except ImportError:
        return names
    locations = [
        (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
        (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
    ]
    for hive, key_name in locations:
        try:
            with winreg.OpenKey(hive, key_name) as key:
                index = 0
                while True:
                    try:
                        value_name, value, _ = winreg.EnumValue(key, index)
                    except OSError:
                        break
                    names.add(re.sub(r"\s*\([^)]*\)\s*$", "", value_name).strip())
                    names.add(Path(str(value)).stem)
                    index += 1
        except OSError:
            continue
    return names


def normalize_font(value: str) -> str:
    return re.sub(r"[\s_-]+", "", value).casefold()


def resolve_font(preferred: str, fallback: str, installed: set[str], warnings: list[str]) -> str:
    normalized = {normalize_font(name): name for name in installed}
    aliases = {
        normalize_font("黑体"): ["SimHei"],
        normalize_font("楷体_GB2312"): ["楷体GB2312", "KaiTi_GB2312", "KaiTi"],
        normalize_font("仿宋_GB2312"): ["仿宋GB2312", "FangSong_GB2312", "FangSong"],
        normalize_font("宋体"): ["SimSun"],
    }
    for candidate in [preferred, fallback, *aliases.get(normalize_font(fallback), [])]:
        if normalize_font(candidate) in normalized:
            if candidate != preferred:
                warnings.append(f"Font fallback used: {preferred} -> {candidate}")
            return candidate
    warnings.append(f"Font unavailable for visual verification: {preferred}; the preferred name was retained")
    return preferred


def set_rfonts(run_or_font, east_asia: str, latin: str) -> None:
    element = getattr(run_or_font, "_element", None)
    if element is None:
        element = getattr(run_or_font, "_parent", None)
    if hasattr(run_or_font, "_element") and run_or_font.__class__.__name__ == "Run":
        r_pr = run_or_font._element.get_or_add_rPr()
    else:
        r_pr = run_or_font._element.get_or_add_rPr() if hasattr(run_or_font, "_element") else None
    if r_pr is None:
        return
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), latin)


def set_east_asia_font_hint(run) -> None:
    r_pr = run._element.get_or_add_rPr()
    hint = r_pr.find(qn("w:hint"))
    if hint is None:
        hint = OxmlElement("w:hint")
        r_pr.append(hint)
    hint.set(qn("w:val"), "eastAsia")


def page_number_rpr(spec: dict, font_name: str):
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("eastAsia", "ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font_name)
    r_pr.append(r_fonts)
    hint = OxmlElement("w:hint")
    hint.set(qn("w:val"), "eastAsia")
    r_pr.append(hint)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    r_pr.append(color)
    bold = OxmlElement("w:b")
    bold.set(qn("w:val"), "1" if spec["bold"] else "0")
    r_pr.append(bold)
    size_value = str(round(spec["size_pt"] * 2))
    for tag in ("w:sz", "w:szCs"):
        size = OxmlElement(tag)
        size.set(qn("w:val"), size_value)
        r_pr.append(size)
    return r_pr


def set_run_black(run) -> None:
    run.font.color.rgb = RGBColor(0, 0, 0)
    color = run._element.get_or_add_rPr().find(qn("w:color"))
    if color is not None:
        color.set(qn("w:val"), "000000")
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attr}"), None)


def set_paragraph_black(paragraph) -> None:
    for run in paragraph.runs:
        set_run_black(run)


def set_paragraph_geometry(paragraph, spec: dict) -> None:
    paragraph.alignment = ALIGNMENTS[spec["alignment"]]
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(spec["space_before_pt"])
    fmt.space_after = Pt(0)
    fmt.widow_control = False
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = p_pr.get_or_add_spacing()
    spacing.set(qn("w:line"), str(round(spec["line_spacing_pt"] * 20)))
    spacing.set(qn("w:lineRule"), "exact")
    spacing.set(qn("w:before"), str(round(spec["space_before_pt"] * 20)))
    spacing.set(qn("w:after"), "0")
    ind = p_pr.get_or_add_ind()
    for attr in ("firstLine", "hanging"):
        key = qn(f"w:{attr}")
        if key in ind.attrib:
            del ind.attrib[key]
    ind.set(qn("w:firstLineChars"), str(round(spec["first_line_chars"] * 100)))


def set_signature_geometry(paragraph, spec: dict, left_twips: int) -> None:
    set_paragraph_geometry(paragraph, spec)
    p_pr = paragraph._p.get_or_add_pPr()
    ind = p_pr.get_or_add_ind()
    ind.set(qn("w:left"), str(left_twips))
    ind.attrib.pop(qn("w:leftChars"), None)
    ind.set(qn("w:firstLine"), "0")
    ind.set(qn("w:firstLineChars"), "0")
    for attr in ("right", "rightChars", "hanging", "hangingChars"):
        ind.attrib.pop(qn(f"w:{attr}"), None)


def ensure_styles(document: Document, profile: dict, fonts: dict[str, str]) -> None:
    for role, style_name in ROLE_NAMES.items():
        spec = signature_style_spec(profile) if role == "signature" else profile["styles"][role]
        font_role = "body" if role == "signature" else role
        try:
            style = document.styles[style_name]
        except KeyError:
            style = document.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = spec["font_latin"]
        style.font.size = Pt(spec["size_pt"])
        style.font.bold = profile["global"]["bold"]
        style.font.color.rgb = RGBColor(0, 0, 0)
        set_rfonts(style, fonts[font_role], spec["font_latin"])
        style.paragraph_format.alignment = ALIGNMENTS[spec["alignment"]]
        style.paragraph_format.space_before = Pt(spec["space_before_pt"])
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.widow_control = False


def apply_role(paragraph, role: str, profile: dict, fonts: dict[str, str], signature_left: int = 0) -> None:
    if role == "skip":
        return
    if role not in ROLE_NAMES:
        raise ValueError(f"Unsupported paragraph role: {role}")
    spec = signature_style_spec(profile) if role == "signature" else profile["styles"][role]
    font_role = "body" if role == "signature" else role
    paragraph.style = ROLE_NAMES[role]
    if role == "signature":
        set_signature_geometry(paragraph, spec, signature_left)
    else:
        set_paragraph_geometry(paragraph, spec)
    for run in paragraph.runs:
        run.font.name = spec["font_latin"]
        run.font.size = Pt(spec["size_pt"])
        run.font.bold = profile["global"]["bold"]
        set_rfonts(run, fonts[font_role], spec["font_latin"])
        set_run_black(run)


def apply_font_only(paragraph, spec: dict, font_name: str, bold: bool | None = None) -> None:
    paragraph.paragraph_format.widow_control = False
    for run in paragraph.runs:
        run.font.name = spec["font_latin"]
        if bold is not None:
            run.font.bold = bold
        set_rfonts(run, font_name, spec["font_latin"])
        set_run_black(run)


def apply_table_format(document: Document, table_specs: list[dict], profile: dict, fonts: dict[str, str]) -> None:
    for table_spec in table_specs:
        title_index = table_spec.get("title_paragraph_index")
        if title_index is not None:
            apply_font_only(
                document.paragraphs[title_index],
                profile["styles"]["main_title"],
                fonts["main_title"],
                profile["global"]["bold"],
            )
        table = document.tables[table_spec["index"]]
        header_rows = table_spec["header_rows"]
        for row_index, row in enumerate(table.rows):
            if row_index == 0 and header_rows >= 1:
                role = "heading1"
            elif row_index == 1 and header_rows >= 2:
                role = "heading2"
            else:
                role = "body"
            spec = profile["styles"][role]
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    apply_font_only(paragraph, spec, fonts[role], profile["global"]["bold"])


def apply_page_setup(document: Document, profile: dict, warnings: list[str]) -> None:
    margins = profile["page"]["margins_cm"]
    for index, section in enumerate(document.sections):
        if section.orientation == WD_ORIENT.LANDSCAPE:
            warnings.append(f"Landscape section {index + 1} was preserved")
            continue
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Cm(margins["top"])
        section.bottom_margin = Cm(margins["bottom"])
        section.left_margin = Cm(margins["left"])
        section.right_margin = Cm(margins["right"])


def clear_paragraph(paragraph) -> None:
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)


def add_page_field(paragraph, spec: dict, font_name: str) -> None:
    clear_paragraph(paragraph)
    paragraph.paragraph_format.widow_control = False
    paragraph.paragraph_format.first_line_indent = None
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    p_pr = paragraph._p.get_or_add_pPr()
    existing_paragraph_r_pr = p_pr.find(qn("w:rPr"))
    if existing_paragraph_r_pr is not None:
        p_pr.remove(existing_paragraph_r_pr)
    p_pr.append(page_number_rpr(spec, font_name))
    left = paragraph.add_run("— ")
    field_begin = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    field_begin._r.append(begin)
    field_instruction = paragraph.add_run()
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE \\* CHARFORMAT "
    field_instruction._r.append(instruction)
    field_separator = paragraph.add_run()
    separator = OxmlElement("w:fldChar")
    separator.set(qn("w:fldCharType"), "separate")
    field_separator._r.append(separator)
    field_result = paragraph.add_run("1")
    field_end = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_end._r.append(end)
    right = paragraph.add_run(" —")
    for run in (left, field_begin, field_instruction, field_separator, field_result, field_end, right):
        run.font.name = font_name
        run.font.size = Pt(spec["size_pt"])
        run.font.bold = spec["bold"]
        set_rfonts(run, font_name, font_name)
        set_east_asia_font_hint(run)
        set_run_black(run)


def set_footer(document: Document, profile: dict, font_name: str) -> None:
    duplex = profile["page"]["print_mode"] == "duplex"
    document.settings.odd_and_even_pages_header_footer = duplex
    page_spec = profile["page_number"]
    for section in document.sections:
        section.different_first_page_header_footer = False
        section.footer.is_linked_to_previous = False
        odd = section.footer.paragraphs[0]
        odd.alignment = WD_ALIGN_PARAGRAPH.RIGHT if duplex else WD_ALIGN_PARAGRAPH.CENTER
        add_page_field(odd, page_spec, font_name)
        if duplex:
            section.even_page_footer.is_linked_to_previous = False
            even = section.even_page_footer.paragraphs[0]
            even.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_page_field(even, page_spec, font_name)


def resolve_fonts(profile: dict) -> tuple[dict[str, str], str, list[str]]:
    installed = installed_font_names()
    warnings: list[str] = []
    fonts = {
        role: resolve_font(spec["font_cn"], spec["font_fallback"], installed, warnings)
        for role, spec in profile["styles"].items()
    }
    page = profile["page_number"]
    page_font = resolve_font(page["font_cn"], page["font_fallback"], installed, warnings)
    return fonts, page_font, warnings


def load_overrides(path: str | None) -> dict | None:
    return read_json(Path(path).expanduser().resolve()) if path else None


def validate_table_specs(items: list[dict], document: Document) -> list[dict]:
    defaults = {item["index"]: item for item in analyze_tables(document)}
    validated = dict(defaults)
    seen: set[int] = set()
    for item in items:
        index = item.get("index")
        if not isinstance(index, int) or index not in defaults:
            raise ValueError(f"Invalid table index: {index}")
        if index in seen:
            raise ValueError(f"Duplicate table index: {index}")
        seen.add(index)
        header_rows = item.get("header_rows")
        if not isinstance(header_rows, int) or header_rows < 0 or header_rows > min(2, len(document.tables[index].rows)):
            raise ValueError(f"Table {index} header_rows must be 0, 1, or 2 and cannot exceed the row count")
        title_index = item.get("title_paragraph_index")
        if title_index is not None and (
            not isinstance(title_index, int) or title_index < 0 or title_index >= len(document.paragraphs)
        ):
            raise ValueError(f"Invalid title paragraph for table {index}: {title_index}")
        validated[index] = {
            **defaults[index],
            "title_paragraph_index": title_index,
            "title_preview": document.paragraphs[title_index].text.strip()[:80] if title_index is not None else None,
            "header_rows": header_rows,
        }
    return [validated[index] for index in sorted(validated)]


def load_format_map(path: str | None, document: Document) -> tuple[list[dict], list[dict]]:
    if not path:
        return classify_paragraphs(document), analyze_tables(document)
    data = read_json(Path(path).expanduser().resolve())
    items = data.get("classifications")
    if not isinstance(items, list):
        raise ValueError("Classification file must contain a classifications array")
    seen: set[int] = set()
    for item in items:
        index = item.get("index")
        role = item.get("role")
        if not isinstance(index, int) or index < 0 or index >= len(document.paragraphs):
            raise ValueError(f"Invalid paragraph index: {index}")
        if role not in {*ROLE_NAMES, "colophon", "skip"}:
            raise ValueError(f"Invalid role at paragraph {index}: {role}")
        if index in seen:
            raise ValueError(f"Duplicate paragraph index: {index}")
        seen.add(index)
    table_items = data.get("tables")
    if table_items is None:
        table_specs = analyze_tables(document)
    elif not isinstance(table_items, list):
        raise ValueError("Classification file tables must be an array")
    else:
        table_specs = validate_table_specs(table_items, document)
    normalize_legacy_roles(document, items)
    return items, table_specs


def _classification_at(classifications: list[dict], index: int) -> dict | None:
    return next((item for item in classifications if item["index"] == index), None)


def _shift_paragraph_indexes(classifications: list[dict], table_specs: list[dict], insertion_index: int) -> None:
    for item in classifications:
        if item["index"] >= insertion_index:
            item["index"] += 1
    for item in table_specs:
        title_index = item.get("title_paragraph_index")
        if title_index is not None and title_index >= insertion_index:
            item["title_paragraph_index"] += 1


def _add_blank_classification(classifications: list[dict], index: int, role: str) -> None:
    item = _classification_at(classifications, index)
    if item is None:
        classifications.append({"index": index, "role": role, "text_preview": ""})
    else:
        item["role"] = role
        item["text_preview"] = ""
    classifications.sort(key=lambda entry: entry["index"])


def _insert_blank_before(document: Document, index: int) -> None:
    paragraph = document.paragraphs[index]
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    Paragraph(element, paragraph._parent)


def normalize_signature_lines(document: Document, classifications: list[dict], table_specs: list[dict]) -> None:
    normalize_legacy_roles(document, classifications)
    for item in list(classifications):
        if item["role"] != "signature":
            continue
        paragraph = document.paragraphs[item["index"]]
        parts = split_signature_text(paragraph.text)
        if parts is None:
            continue
        # Do not flatten fields, links, drawings, bookmarks or section boundaries.
        p_pr = paragraph._p.pPr
        if p_pr is not None and p_pr.find(qn("w:sectPr")) is not None:
            raise ValueError("Signature with a section break cannot be split automatically")
        for child in paragraph._p:
            if child.tag == qn("w:pPr"):
                continue
            if child.tag != qn("w:r") or any(
                node.tag not in {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br")} for node in child
            ):
                raise ValueError("Signature contains complex content; provide separate plain office/date paragraphs")
            if any(node.tag == qn("w:br") and node.get(qn("w:type")) in {"page", "column"} for node in child):
                raise ValueError("Signature contains a page or column break; split it manually")
        office, date = parts
        paragraph.text = office
        item["text_preview"] = office[:80]
        insertion_index = item["index"] + 1
        _shift_paragraph_indexes(classifications, table_specs, insertion_index)
        element = OxmlElement("w:p")
        paragraph._p.addnext(element)
        Paragraph(element, paragraph._parent).text = date
        classifications.append({"index": insertion_index, "role": "signature", "text_preview": date[:80]})
    classifications.sort(key=lambda entry: entry["index"])


def _remove_spacing(document: Document, classifications: list[dict], table_specs: list[dict], index: int) -> None:
    paragraph = document.paragraphs[index]
    if not is_spacing_paragraph(paragraph):
        raise ValueError("Cannot remove a non-spacing paragraph")
    if any(item.get("title_paragraph_index") == index for item in table_specs):
        raise ValueError("A table title cannot be removed as spacing")
    paragraph._p.getparent().remove(paragraph._p)
    classifications[:] = [item for item in classifications if item["index"] != index]
    for item in classifications:
        if item["index"] > index:
            item["index"] -= 1
    for item in table_specs:
        if item.get("title_paragraph_index") is not None and item["title_paragraph_index"] > index:
            item["title_paragraph_index"] -= 1


def ensure_required_spacing(document: Document, classifications: list[dict], table_specs: list[dict]) -> None:
    for title_item in [item for item in classifications if item["role"] == "main_title"]:
        paragraphs = document.paragraphs
        roles = {item["index"]: item["role"] for item in classifications}
        end_index, blanks = title_block_layout(document, roles, title_item["index"])
        end_element = paragraphs[end_index]._p
        # Keep element identities so moving legacy spacing also remaps table titles
        # and confirmed classifications, without reclassifying the user's content.
        entries = [(item, paragraphs[item["index"]]._p) for item in classifications]
        table_titles = [
            (item, paragraphs[item["title_paragraph_index"]]._p)
            for item in table_specs if item.get("title_paragraph_index") is not None
        ]
        if blanks:
            keep_index = next((index for index in blanks if index > end_index), blanks[0])
            spacer = paragraphs[keep_index]._p
            for index in blanks:
                if index != keep_index:
                    element = paragraphs[index]._p
                    element.getparent().remove(element)
            if end_element.getnext() is not spacer:
                end_element.addnext(spacer)
        else:
            spacer = OxmlElement("w:p")
            end_element.addnext(spacer)
        indexes = {paragraph._p: index for index, paragraph in enumerate(document.paragraphs)}
        classifications[:] = [item for item, element in entries if element in indexes]
        for item, element in entries:
            if element in indexes:
                item["index"] = indexes[element]
        for item, element in table_titles:
            item["title_paragraph_index"] = indexes[element]
        _add_blank_classification(classifications, indexes[spacer], "body")

    signature_items = [
        item for item in classifications
        if item["role"] == "signature" and document.paragraphs[item["index"]].text.strip()
    ]
    if signature_items:
        first_index = min(item["index"] for item in signature_items)
        # Collapse only plain spacing immediately above the signature, keeping
        # one paragraph; never remove an image or a page/section break.
        while first_index >= 2:
            previous = document.paragraphs[first_index - 1]
            earlier = document.paragraphs[first_index - 2]
            if (
                earlier._p.getnext() is not previous._p
                or previous._p.getnext() is not document.paragraphs[first_index]._p
                or not is_spacing_paragraph(previous) or not is_spacing_paragraph(earlier)
            ):
                break
            _remove_spacing(document, classifications, table_specs, first_index - 2)
            first_index -= 1
        first_paragraph = document.paragraphs[first_index]
        preceding = first_paragraph._p.getprevious()
        if preceding is not None and preceding.tag == qn("w:p") and is_spacing_paragraph(Paragraph(preceding, first_paragraph._parent)):
            _add_blank_classification(classifications, first_index - 1, "signature")
        else:
            _shift_paragraph_indexes(classifications, table_specs, first_index)
            _insert_blank_before(document, first_index)
            _add_blank_classification(classifications, first_index, "signature")


def apply_document_format(document: Document, classifications: list[dict], table_specs: list[dict], profile: dict) -> list[str]:
    fonts, page_font, warnings = resolve_fonts(profile)
    apply_page_setup(document, profile, warnings)
    ensure_styles(document, profile, fonts)
    signature_indents = signature_left_indents(document, profile, {item["index"]: item["role"] for item in classifications})
    if any(indent == 0 for indent in signature_indents.values()):
        warnings.append("Long signature uses the full text width and may wrap; review the rendered layout")
    table_title_indexes = {
        item["title_paragraph_index"] for item in table_specs if item.get("title_paragraph_index") is not None
    }
    for item in classifications:
        role = "skip" if item["index"] in table_title_indexes else item["role"]
        apply_role(document.paragraphs[item["index"]], role, profile, fonts, signature_indents.get(item["index"], 0))
    apply_table_format(document, table_specs, profile, fonts)
    set_footer(document, profile, page_font)
    return warnings


def read_source_text(args: argparse.Namespace) -> str:
    if args.text is not None:
        return args.text
    return Path(args.input).expanduser().resolve().read_text(encoding="utf-8")


def new_output_path(output: str | None) -> Path:
    if output:
        result = Path(output).expanduser().resolve()
        if result.suffix.lower() != ".docx":
            raise ValueError("Output must use the .docx extension")
        return unique_output_path(result)
    return unique_output_path(Path.cwd() / "通用公文.docx")


def add_text(document: Document, text: str) -> list[dict]:
    explicit: dict[int, str] = {}
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        role = None
        if line.startswith("### "):
            role, line = "heading2", line[4:].strip()
        elif line.startswith("## "):
            role, line = "heading1", line[3:].strip()
        elif line.startswith("# "):
            role, line = "main_title", line[2:].strip()
        elif line.startswith("[说明]"):
            role, line = "description", line[4:].strip()
        elif line.startswith("[落款]"):
            role, line = "signature", line[4:].strip()
        elif line.startswith("[版记]"):
            role, line = "colophon", line[4:].strip()
        paragraph = document.add_paragraph(line)
        if role:
            explicit[len(document.paragraphs) - 1] = role
    classifications = classify_paragraphs(document)
    for item in classifications:
        if item["index"] in explicit:
            item["role"] = explicit[item["index"]]
    normalize_legacy_roles(document, classifications)
    return classifications


def finalize(document: Document, output: Path, profile: dict, classifications: list[dict], table_specs: list[dict]) -> dict:
    output.parent.mkdir(parents=True, exist_ok=True)
    normalize_signature_lines(document, classifications, table_specs)
    ensure_required_spacing(document, classifications, table_specs)
    warnings = apply_document_format(document, classifications, table_specs, profile)
    document.save(output)
    scrub_docx(output)
    try:
        validation = validate_document(output, profile, classifications, table_specs)
    except Exception:
        output.unlink(missing_ok=True)
        raise
    if not validation["valid"]:
        output.unlink(missing_ok=True)
        raise RuntimeError("Generated DOCX failed structural validation: " + "; ".join(validation["errors"]))
    return {"output": str(output), "warnings": sorted(set(warnings + validation["warnings"])), "validation": "structural_pass"}


def command_create(args: argparse.Namespace) -> None:
    profile = active_profile(load_overrides(args.overrides))
    if args.print_mode:
        profile["page"]["print_mode"] = args.print_mode
    document = Document()
    classifications = add_text(document, read_source_text(args))
    result = finalize(document, new_output_path(args.output), profile, classifications, [])
    print(json.dumps(result, ensure_ascii=False, indent=2))


def command_format_existing(args: argparse.Namespace) -> None:
    if not args.confirmed:
        raise SystemExit("Formatting requires --confirmed after the user reviews the inspection summary")
    source = Path(args.input).expanduser().resolve()
    inspection = inspect_document(source)
    if inspection.get("blockers"):
        raise SystemExit("; ".join(inspection["blockers"]))
    profile = active_profile(load_overrides(args.overrides))
    if args.print_mode:
        profile["page"]["print_mode"] = args.print_mode
    document = Document(source)
    classifications, table_specs = load_format_map(args.classification, document)
    output = choose_output_path(source, args.output)
    result = finalize(document, output, profile, classifications, table_specs)
    result["warnings"] = sorted(set(result["warnings"] + inspection.get("warnings", [])))
    print(json.dumps(result, ensure_ascii=False, indent=2))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create or reformat a general official-document DOCX")
    sub = parser.add_subparsers(dest="command", required=True)
    create = sub.add_parser("create")
    source = create.add_mutually_exclusive_group(required=True)
    source.add_argument("--input", help="UTF-8 text or Markdown file")
    source.add_argument("--text", help="Inline plain text or Markdown")
    create.add_argument("--output")
    create.add_argument("--overrides", help="Temporary JSON override file")
    create.add_argument("--print-mode", choices=["single", "duplex"])
    create.set_defaults(func=command_create)

    existing = sub.add_parser("format-existing")
    existing.add_argument("--input", required=True)
    existing.add_argument("--output")
    existing.add_argument("--classification", help="Reviewed classification JSON")
    existing.add_argument("--overrides", help="Temporary JSON override file")
    existing.add_argument("--print-mode", choices=["single", "duplex"])
    existing.add_argument("--confirmed", action="store_true")
    existing.set_defaults(func=command_format_existing)
    return parser


if __name__ == "__main__":
    configure_utf8_stdio()
    arguments = build_parser().parse_args()
    try:
        arguments.func(arguments)
    except (KeyError, TypeError, ValueError) as exc:
        raise SystemExit(str(exc)) from exc
