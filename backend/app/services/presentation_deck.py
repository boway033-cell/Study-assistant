"""Evidence-grounded Chinese literature-report deck generation.

The pipeline adapts nature-paper2ppt's paper-type routing and nature-polishing's
claim-evidence-boundary discipline.  It never asks the model to invent missing
results: each slide must retain local chunk/page source IDs.
"""
from __future__ import annotations

import json
import math
import re
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy import select

from backend.app.core.config import settings
from backend.app.core.database import SessionLocal
from backend.app.models import Book, Chapter, Chunk, LiteratureResource, PaperProfile, PresentationDeck
from backend.app.services.llm import LLMRouter, load_llm_config, parse_json_response
from backend.app.worker.tasks import update_progress


PAPER_TYPE_LABELS = {
    "discovery": "发现 / 机制",
    "methods": "方法 / 算法",
    "resource": "资源 / 数据集",
    "clinical": "临床 / 人群",
    "materials": "材料 / 工程",
    "review": "综述 / 观点",
}

TYPE_KEYWORDS = {
    "clinical": "trial cohort patient clinical randomized intervention survival meta-analysis 临床 患者 队列 试验",
    "methods": "method algorithm model benchmark framework architecture dataset baseline 方法 算法 模型 基准 架构",
    "resource": "atlas resource database cohort dataset accession repository 图谱 资源 数据库 数据集",
    "materials": "material synthesis fabrication device performance xrd sem catalyst 材料 合成 器件 表征 性能",
    "review": "review perspective commentary systematic review 综述 述评 展望",
    "discovery": "mechanism pathway phenotype cell gene protein discovery 机制 通路 表型 细胞 基因",
}


def classify_paper_type(title: str, text: str) -> str:
    hay = f"{title}\n{text[:18000]}".lower()
    scores = {kind: sum(hay.count(w) for w in words.split()) for kind, words in TYPE_KEYWORDS.items()}
    return max(scores, key=scores.get) if max(scores.values(), default=0) else "discovery"


def _descendants(chapters: list[Chapter], root_ids: set[int]) -> set[int]:
    children: dict[int | None, list[int]] = {}
    for chapter in chapters:
        children.setdefault(chapter.parent_id, []).append(chapter.id)
    result = set(root_ids); stack = list(root_ids)
    while stack:
        child_ids = children.get(stack.pop(), [])
        result.update(child_ids); stack.extend(child_ids)
    return result


def _scope_root(chapter_id: int | None, parent_map: dict[int, int | None], roots: set[int]) -> int | None:
    current = chapter_id
    seen = set()
    while current is not None and current not in seen:
        if current in roots:
            return current
        seen.add(current); current = parent_map.get(current)
    return None


def _sample_group(chunks: list[Chunk], quota: int) -> tuple[list[tuple[Chunk, str]], int]:
    """均匀覆盖首/中/尾，避免只读取章节前部。"""
    usable = [(chunk, (chunk.content or "").strip()) for chunk in chunks if (chunk.content or "").strip()]
    if not usable or quota <= 0:
        return [], 0
    selected: list[tuple[Chunk, str]] = []
    remaining = quota
    pending = list(range(len(usable)))
    # 先按均匀位置各取一轮，再利用剩余额度补充相邻内容。
    target_count = min(len(usable), max(3, math.ceil(quota / 2800)))
    if target_count == 1:
        indexes = [0]
    else:
        indexes = sorted({round(i * (len(usable) - 1) / (target_count - 1)) for i in range(target_count)})
    indexes += [i for i in pending if i not in indexes]
    for index in indexes:
        chunk, content = usable[index]
        if remaining <= 0:
            break
        excerpt = content[:min(4000, remaining)]
        if excerpt:
            selected.append((chunk, excerpt)); remaining -= len(excerpt)
    selected.sort(key=lambda item: item[0].chunk_index)
    return selected, quota - remaining


def collect_selection(db, book_id: int, chapter_ids: list[int], chunk_ids: list[int], selected_text: str,
                      resource_ids: list[int] | None = None, max_source_chars: int = 52000) -> dict:
    book = db.get(Book, book_id)
    if not book:
        raise ValueError("文献不存在")
    chapters = list(db.scalars(select(Chapter).where(Chapter.book_id == book_id)
                               .order_by(Chapter.order_index)).all())
    chapter_map = {chapter.id: chapter for chapter in chapters}
    parent_map = {chapter.id: chapter.parent_id for chapter in chapters}
    query = select(Chunk).where(Chunk.book_id == book_id)
    if chunk_ids:
        query = query.where(Chunk.id.in_(chunk_ids))
    elif chapter_ids:
        query = query.where(Chunk.chapter_id.in_(_descendants(chapters, set(chapter_ids))))
    main_chunks = list(db.scalars(query.order_by(Chunk.chunk_index)).all())

    groups: list[dict] = []
    if chunk_ids:
        groups.append({"key": "explicit_chunks", "title": "指定片段", "book_id": book_id,
                       "resource_id": None, "chunks": main_chunks})
    else:
        roots = set(chapter_ids) if chapter_ids else {c.id for c in chapters if c.parent_id is None or c.level == 1}
        if not roots:
            groups.append({"key": f"book:{book_id}", "title": book.title, "book_id": book_id,
                           "resource_id": None, "chunks": main_chunks})
        else:
            bucket = {root: [] for root in roots}
            other = []
            for chunk in main_chunks:
                root = _scope_root(chunk.chapter_id, parent_map, roots)
                (bucket[root] if root in bucket else other).append(chunk)
            for root in sorted(roots, key=lambda rid: chapter_map.get(rid).order_index if rid in chapter_map else rid):
                groups.append({"key": f"chapter:{root}", "title": chapter_map[root].title if root in chapter_map else f"章节 {root}",
                               "book_id": book_id, "resource_id": None, "chunks": bucket[root]})
            if other:
                groups.append({"key": "unmapped", "title": "未映射章节", "book_id": book_id,
                               "resource_id": None, "chunks": other})

    selected_resources = []
    if resource_ids:
        selected_resources = list(db.scalars(select(LiteratureResource).where(
            LiteratureResource.book_id == book_id, LiteratureResource.id.in_(set(resource_ids)))).all())
        if len(selected_resources) != len(set(resource_ids)):
            raise ValueError("部分补充材料不存在或不属于当前文献")
        for resource in selected_resources:
            if resource.resource_book_id:
                resource_chunks = list(db.scalars(select(Chunk).where(Chunk.book_id == resource.resource_book_id)
                                                  .order_by(Chunk.chunk_index)).all())
                groups.append({"key": f"resource:{resource.id}", "title": resource.title,
                               "book_id": resource.resource_book_id, "resource_id": resource.id,
                               "chunks": resource_chunks})

    user_text = selected_text.strip()[:12000]
    source_budget = max(8000, min(max_source_chars, 100000)) - len(user_text)
    available_by_group = [sum(len((c.content or "").strip()) for c in group["chunks"]) for group in groups]
    weights = [math.sqrt(max(chars, 1)) for chars in available_by_group]
    weight_total = sum(weights) or 1
    quotas = [max(800, int(source_budget * weight / weight_total)) for weight in weights]
    if sum(quotas) > source_budget and quotas:
        scale = source_budget / sum(quotas)
        quotas = [max(300, int(quota * scale)) for quota in quotas]

    sources = []
    coverage_groups = []
    for group, available_chars, quota in zip(groups, available_by_group, quotas):
        sampled, sampled_chars = _sample_group(group["chunks"], quota)
        for chunk, content in sampled:
            sources.append({"source_id": f"chunk:{chunk.id}", "chunk_id": chunk.id,
                            "book_id": chunk.book_id, "resource_id": group["resource_id"],
                            "chapter_id": chunk.chapter_id, "chapter_title": chapter_map.get(chunk.chapter_id).title if chunk.chapter_id in chapter_map else None,
                            "page_start": chunk.page_start, "page_end": chunk.page_end, "text": content})
        coverage_groups.append({"key": group["key"], "title": group["title"],
                                "available_chars": available_chars, "sampled_chars": sampled_chars,
                                "total_chunks": len(group["chunks"]), "sampled_chunks": len(sampled),
                                "coverage": round(sampled_chars / available_chars, 4) if available_chars else 0})
    if user_text:
        sources.insert(0, {"source_id": "selection:user", "chunk_id": None, "book_id": book_id,
                           "resource_id": None, "chapter_id": None, "chapter_title": None,
                           "page_start": None, "page_end": None, "text": user_text})
    if not sources:
        raise ValueError("所选范围没有可用于汇报的正文")
    total_available = sum(available_by_group)
    total_sampled = sum(group["sampled_chars"] for group in coverage_groups)
    covered_groups = sum(1 for group in coverage_groups if group["sampled_chunks"])
    return {
        "book": {"id": book.id, "title": book.title},
        "chapter_ids": chapter_ids, "chunk_ids": chunk_ids,
        "resource_ids": resource_ids or [],
        "chapter_titles": [chapter_map[i].title for i in chapter_ids if i in chapter_map],
        "sources": sources, "coverage": {"max_source_chars": max_source_chars,
            "available_chars": total_available, "sampled_chars": total_sampled,
            "content_coverage": round(total_sampled / total_available, 4) if total_available else 0,
            "structure_coverage": round(covered_groups / len(coverage_groups), 4) if coverage_groups else 0,
            "covered_groups": covered_groups, "total_groups": len(coverage_groups),
            "unlocated_user_selection": bool(user_text), "groups": coverage_groups},
    }


def _local_outline(title: str, paper_type: str, sources: list[dict], slide_count: int) -> list[dict]:
    arcs = {
        "methods": ["研究背景与当前瓶颈", "核心问题", "方法总览", "关键设计", "评测设置", "主要结果", "稳健性与失败案例", "适用边界", "总结"],
        "resource": ["研究背景", "资源概览", "样本与数据设计", "生成与质控流程", "主要图谱", "关键洞见", "验证与复用", "局限性", "总结"],
        "clinical": ["临床问题", "研究问题", "研究设计", "终点与变量", "主要结果", "次要分析", "偏倚与局限", "实践意义", "总结"],
        "materials": ["技术挑战", "设计原理", "制备与装置", "关键表征", "性能证据", "构效关系", "稳定性与边界", "局限性", "总结"],
        "review": ["主题为何重要", "概念框架", "主题一", "主题二", "主题三", "争议与缺口", "作者综合观点", "未来方向", "总结"],
        "discovery": ["研究背景", "知识缺口", "核心问题与假设", "实验设计", "关键证据一", "关键证据二", "验证与稳健性", "机制模型", "局限性", "总结"],
    }
    wanted = max(6, min(slide_count, 18))
    names = arcs[paper_type]
    if wanted - 1 < len(names):
        names = names[: wanted - 2] + [names[-1]]
    while len(names) < wanted - 1:
        names.insert(-1, f"补充证据 {len(names) - 3}")
    slides = [{"title": title, "kind": "cover", "claim": f"{PAPER_TYPE_LABELS[paper_type]}论文中文汇报",
               "bullets": [], "source_ids": []}]
    for i, name in enumerate(names):
        src = sources[min(i, len(sources) - 1)]
        sentence = re.split(r"(?<=[。！？.!?])\s*|\n+", src["text"])[0].strip()[:130]
        slides.append({"title": name, "kind": "content", "claim": sentence or "请结合原文核对本页论点",
                       "bullets": ["本页由本地证据片段生成；建议在汇报前核对原图与数值。"],
                       "source_ids": [src["source_id"]]})
    return slides[:wanted]


async def _ai_outline(provider, title: str, paper_type: str, sources: list[dict], options: dict) -> list[dict] | None:
    source_text = "\n\n".join(
        f"[{s['source_id']}; p.{s['page_start'] or '?'}-{s['page_end'] or s['page_start'] or '?'}]\n{s['text']}"
        for s in sources
    )
    prompt = f"""你是严谨的中文学术汇报编辑。论文类型：{PAPER_TYPE_LABELS[paper_type]}。
受众：{options['audience']}；目的：{options['purpose']}；目标 {options['slide_count']} 页，{options['duration_minutes']} 分钟。
按“为什么重要→知识缺口→作者做了什么→关键证据→可信度→意义/复用→边界”建立证据链。
只使用下方来源。不得补造数字、因果关系、实验或结论；不确定就写“原文未说明”。
每页只表达一个结论，中文短句，术语和数值逐字忠于来源。每个事实页必须列 source_ids。
输出严格 JSON 数组，每项字段：title, kind(cover/content/evidence/limitations/summary), claim,
bullets(0-4条，每条不超过45字), source_ids(只能引用给定ID)。封面可无来源。不要 markdown。

来源：
{source_text}"""
    chunks = []
    async for delta in provider.stream_chat([{"role": "system", "content": "遵循 claim-evidence-boundary；只输出 JSON。"},
                                              {"role": "user", "content": prompt}]):
        chunks.append(delta)
    parsed = parse_json_response("".join(chunks))
    if not isinstance(parsed, list):
        return None
    allowed = {s["source_id"] for s in sources}
    result = []
    for item in parsed[:18]:
        if not isinstance(item, dict):
            continue
        ids = [x for x in item.get("source_ids", []) if x in allowed]
        result.append({"title": str(item.get("title") or "未命名页面")[:70],
                       "kind": str(item.get("kind") or "content")[:20],
                       "claim": str(item.get("claim") or "")[:220],
                       "bullets": [str(x)[:100] for x in item.get("bullets", [])[:4]],
                       "source_ids": ids})
    return result if len(result) >= 5 else None


def _terms(text: str) -> set[str]:
    latin = {x.lower() for x in re.findall(r"[A-Za-z][A-Za-z0-9_.+-]{2,}", text)}
    chinese = re.findall(r"[\u4e00-\u9fff]{2,}", text)
    bigrams = {word[i:i + 2] for word in chinese for i in range(len(word) - 1)}
    return latin | bigrams


def audit_claim_sources(outline: list[dict], sources: list[dict]) -> dict:
    """保守的本地一致性审计：定位、数字和术语；不把词面相似冒充语义蕴含。"""
    source_map = {source["source_id"]: source for source in sources}
    results = []
    counts = {"supported": 0, "partial": 0, "needs_review": 0, "unsupported": 0, "unlocated": 0}
    blocking = []
    for index, slide in enumerate(outline, 1):
        if index == 1 or slide.get("kind") == "cover":
            continue
        source_ids = list(dict.fromkeys(slide.get("source_ids") or []))
        valid = [source_map[source_id] for source_id in source_ids if source_id in source_map]
        invalid_ids = [source_id for source_id in source_ids if source_id not in source_map]
        claim = " ".join([str(slide.get("claim") or ""), *[str(x) for x in slide.get("bullets", [])]])
        evidence = "\n".join(source.get("text") or "" for source in valid)
        claim_numbers = set(re.findall(r"(?<!\w)\d+(?:\.\d+)?%?", claim))
        source_numbers = set(re.findall(r"(?<!\w)\d+(?:\.\d+)?%?", evidence))
        missing_numbers = sorted(claim_numbers - source_numbers)
        claim_terms = _terms(claim); source_terms = _terms(evidence)
        overlap = len(claim_terms & source_terms) / max(len(claim_terms), 1)
        located = any(source.get("page_start") or source.get("chapter_id") for source in valid)
        issues = []
        if invalid_ids: issues.append(f"无效来源：{', '.join(invalid_ids)}")
        if not valid: issues.append("没有可核对的来源")
        if missing_numbers: issues.append("来源中未找到数字：" + "、".join(missing_numbers[:8]))
        if valid and overlap < .08: issues.append("主张与来源词面重合度低，需人工核对语义")
        if not located: issues.append("来源缺少页码或章节定位")
        if not valid or invalid_ids or missing_numbers:
            status = "unsupported"; blocking.append(index)
        elif not located:
            status = "unlocated"
        elif overlap < .08:
            status = "needs_review"
        elif overlap < .18:
            status = "partial"
        else:
            status = "supported"
        counts[status] += 1
        results.append({"slide": index, "title": slide.get("title"), "status": status,
                        "source_ids": source_ids, "overlap_score": round(overlap, 4),
                        "missing_numbers": missing_numbers, "issues": issues})
    total = len(results)
    accepted = counts["supported"] + counts["partial"]
    return {"ok": not blocking, "blocking_slides": blocking, "counts": counts,
            "accepted_ratio": round(accepted / total, 4) if total else 1.0,
            "items": results,
            "boundary": "本地审计只验证定位、数字和术语一致性；needs_review 仍需人工判断语义是否成立。"}


def _add_textbox(slide, left, top, width, height, text, size, color, bold=False, font="Microsoft YaHei"):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.TOP
    p = frame.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.LEFT
    p.font.name = font; p.font.size = size; p.font.bold = bold; p.font.color.rgb = RGBColor(*color)
    return box


def _extract_page_figure(pdf_path: Path, page_no: int, output: Path) -> Path | None:
    """Extract the largest embedded raster on a source page without cropping labels."""
    if not pdf_path.exists() or page_no < 1:
        return None
    try:
        import fitz
        doc = fitz.open(pdf_path)
        try:
            if page_no > len(doc): return None
            images = doc[page_no - 1].get_images(full=True)
            candidates = []
            for info in images:
                width, height = int(info[2]), int(info[3])
                if width >= 420 and height >= 260:
                    candidates.append((width * height, info[0]))
            if not candidates: return None
            pix = fitz.Pixmap(doc, max(candidates)[1])
            if pix.n - pix.alpha > 3:
                pix = fitz.Pixmap(fitz.csRGB, pix)
            pix.save(output)
            return output
        finally:
            doc.close()
    except Exception:
        return None


def render_pptx(path: Path, book: Book, profile: PaperProfile | None, outline: list[dict], sources: list[dict], paper_type: str, options: dict) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    source_map = {s["source_id"]: s for s in sources}
    pdf_path = settings.uploads_dir / book.file_path
    with tempfile.TemporaryDirectory(prefix="deck_fig_") as tmp:
      tmp_dir = Path(tmp)
      for idx, item in enumerate(outline):
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(247, 244, 237)
        if idx == 0:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.3), prs.slide_height)
            shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(24, 54, 56); shape.line.fill.background()
            _add_textbox(slide, Inches(.72), Inches(.75), Inches(2.8), Inches(.35), "LITERATURE REPORT", Pt(15), (210,181,143), True)
            _add_textbox(slide, Inches(4.85), Inches(1.25), Inches(7.45), Inches(2.2), item["title"], Pt(36), (35,47,45), True)
            _add_textbox(slide, Inches(4.9), Inches(3.8), Inches(6.8), Inches(.8), item.get("claim", ""), Pt(20), (92,87,76))
            meta = " · ".join(x for x in [profile.authors if profile else None, profile.journal if profile else None,
                                            str(profile.published_year) if profile and profile.published_year else None] if x)
            _add_textbox(slide, Inches(4.9), Inches(5.75), Inches(6.8), Inches(.45), meta or "本地知识库文献", Pt(13), (115,108,95))
        else:
            _add_textbox(slide, Inches(.72), Inches(.42), Inches(11.6), Inches(.65), item["title"], Pt(28), (28,54,54), True)
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.72), Inches(1.22), Inches(.8), Inches(.06))
            accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(182,128,76); accent.line.fill.background()
            hero = None
            if options.get("include_figures") and book.file_type == "pdf":
                for sid in item.get("source_ids", []):
                    src = source_map.get(sid, {}); page_no = src.get("page_start")
                    if page_no:
                        hero = _extract_page_figure(pdf_path, int(page_no), tmp_dir / f"{idx}-{page_no}.png")
                        if hero: break
            text_width = Inches(5.65) if hero else Inches(11.45)
            _add_textbox(slide, Inches(.86), Inches(1.62), text_width, Inches(1.3), item.get("claim", ""), Pt(24), (39,42,39), True)
            bullet_text = "\n".join(f"• {x}" for x in item.get("bullets", []))
            _add_textbox(slide, Inches(1.02), Inches(3.15), Inches(5.35) if hero else Inches(10.9), Inches(2.55), bullet_text, Pt(18), (69,70,64))
            if hero:
                from PIL import Image
                with Image.open(hero) as im: ratio = im.width / max(im.height, 1)
                max_w, max_h = Inches(5.5), Inches(4.7)
                pic_w = min(max_w, int(max_h * ratio)); pic_h = int(pic_w / ratio)
                slide.shapes.add_picture(str(hero), Inches(7.15) + (max_w - pic_w) // 2,
                                         Inches(1.48) + (max_h - pic_h) // 2, width=pic_w, height=pic_h)
            refs = []
            for sid in item.get("source_ids", []):
                src = source_map.get(sid, {})
                page = src.get("page_start")
                refs.append(f"{sid}" + (f" · p.{page}" if page else ""))
            _add_textbox(slide, Inches(.76), Inches(6.78), Inches(11.8), Inches(.28), "来源：" + "；".join(refs), Pt(9), (125,119,108))
        _add_textbox(slide, Inches(12.25), Inches(6.84), Inches(.45), Inches(.25), f"{idx+1:02d}", Pt(9), (125,119,108), True)
        notes = [f"[Deck] {book.title}", f"[PaperType] {paper_type}", "[Sources]"]
        for sid in item.get("source_ids", []):
            src = source_map.get(sid, {})
            notes.append(f"- {sid}; pages={src.get('page_start')}-{src.get('page_end')}")
        try:
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)
        except Exception:
            pass
    path.parent.mkdir(parents=True, exist_ok=True); prs.save(path)


def audit_pptx(path: Path, outline: list[dict]) -> dict:
    from pptx import Presentation
    prs = Presentation(path)
    issues = []
    if len(prs.slides) != len(outline): issues.append("幻灯片数量与提纲不一致")
    for i, (slide, spec) in enumerate(zip(prs.slides, outline), 1):
        text = " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text_frame"))
        if i > 1 and not spec.get("source_ids"): issues.append(f"第{i}页缺少来源定位")
        if len(text) > 520: issues.append(f"第{i}页文字偏多")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                issues.append(f"第{i}页存在越界对象"); break
    return {"ok": not issues, "slide_count": len(prs.slides), "issues": issues}


def validate_outline(outline: list[dict], sources: list[dict]) -> list[dict]:
    if not 2 <= len(outline) <= 24:
        raise ValueError("提纲页数需为 2–24 页")
    allowed = {source["source_id"] for source in sources}
    cleaned = []
    for index, raw in enumerate(outline):
        if not isinstance(raw, dict):
            raise ValueError(f"第 {index + 1} 页格式无效")
        source_ids = list(dict.fromkeys(str(item) for item in raw.get("source_ids", [])))
        invalid = [source_id for source_id in source_ids if source_id not in allowed]
        if invalid:
            raise ValueError(f"第 {index + 1} 页包含无效来源：{', '.join(invalid)}")
        title = str(raw.get("title") or "").strip()[:70]
        if not title:
            raise ValueError(f"第 {index + 1} 页缺少标题")
        kind = "cover" if index == 0 else str(raw.get("kind") or "content")[:20]
        if index > 0 and not source_ids:
            raise ValueError(f"第 {index + 1} 页至少需要一个来源")
        cleaned.append({"title": title, "kind": kind,
                        "claim": str(raw.get("claim") or "").strip()[:220],
                        "bullets": [str(item).strip()[:100] for item in raw.get("bullets", [])[:4] if str(item).strip()],
                        "source_ids": source_ids})
    return cleaned


def _rights_policy(db, book_id: int, options: dict) -> dict:
    resources = list(db.scalars(select(LiteratureResource).where(LiteratureResource.book_id == book_id)).all())
    reusable = [resource for resource in resources if resource.allow_reuse or resource.rights_status in
                {"open_license", "permission_granted", "public_domain"}]
    unknown = [resource for resource in resources if resource.rights_status in {"not_evaluated", "undetermined", "in_copyright", "restricted"}]
    scope = options.get("use_scope", "personal")
    requested = bool(options.get("include_figures"))
    acknowledged = bool(options.get("rights_acknowledged"))
    issues = []
    allow_figures = requested
    if requested and scope in {"public", "commercial"} and not reusable:
        allow_figures = False; issues.append("公开或商业用途没有可复用许可，已自动关闭原图嵌入")
    elif requested and unknown and not acknowledged:
        allow_figures = False; issues.append("存在未评估权利资源且尚未确认，已自动关闭原图嵌入")
    manifest = [{"id": resource.id, "role": resource.role, "title": resource.title,
                 "license_expression": resource.license_expression,
                 "rights_statement_uri": resource.rights_statement_uri,
                 "rights_status": resource.rights_status, "allow_reuse": bool(resource.allow_reuse),
                 "attribution": resource.attribution} for resource in resources]
    return {"use_scope": scope, "requested_figures": requested, "allow_figures": allow_figures,
            "acknowledged": acknowledged, "issues": issues, "resources": manifest,
            "boundary": "系统记录权利元数据并执行保守复用策略，不替代法律判断。"}


async def generate_outline_task(record, deck_id: int) -> dict:
    db = SessionLocal()
    try:
        deck = db.get(PresentationDeck, deck_id)
        if not deck:
            raise ValueError("汇报任务不存在")
        deck.status = "outlining"; db.commit()
        selection = json.loads(deck.selection_json or "{}"); options = json.loads(deck.options_json or "{}")
        book = db.get(Book, deck.book_id); sources = selection["sources"]
        paper_type = classify_paper_type(book.title, "\n".join(source["text"] for source in sources))
        deck.paper_type = paper_type
        update_progress(record, .18, "deck_outline", f"已识别为{PAPER_TYPE_LABELS[paper_type]}论文")
        outline = None; cfg = load_llm_config(db)
        if cfg.get("deepseek_api_key"):
            update_progress(record, .48, "deck_outline", "AI 正在按证据链组织可编辑提纲")
            outline = await _ai_outline(LLMRouter.get("auto", cfg), book.title, paper_type, sources, options)
        if not outline:
            outline = _local_outline(book.title, paper_type, sources, options["slide_count"])
        outline = validate_outline(outline, sources)
        claim_audit = audit_claim_sources(outline, sources)
        deck.outline_json = json.dumps(outline, ensure_ascii=False)
        deck.qa_json = json.dumps({"stage": "outline", "claim_source": claim_audit,
                                   "coverage": selection.get("coverage", {})}, ensure_ascii=False)
        deck.status = "outline_ready"; db.commit()
        update_progress(record, 1, "deck_outline", "提纲已生成，等待人工确认")
        return {"deck_id": deck.id, "status": deck.status, "slide_count": len(outline),
                "coverage": selection.get("coverage", {}), "claim_source": claim_audit}
    except Exception as exc:
        db.rollback(); deck = db.get(PresentationDeck, deck_id)
        if deck: deck.status = "failed"; deck.error_msg = str(exc); db.commit()
        raise
    finally:
        db.close()


async def render_deck_task(record, deck_id: int) -> dict:
    from backend.app.services.powerpoint_render import render_powerpoint_preview

    db = SessionLocal()
    try:
        deck = db.get(PresentationDeck, deck_id)
        if not deck:
            raise ValueError("汇报任务不存在")
        selection = json.loads(deck.selection_json or "{}"); options = json.loads(deck.options_json or "{}")
        outline = validate_outline(json.loads(deck.outline_json or "[]"), selection.get("sources", []))
        book = db.get(Book, deck.book_id); profile = db.get(PaperProfile, deck.book_id)
        rights = _rights_policy(db, deck.book_id, options)
        render_options = {**options, "include_figures": rights["allow_figures"]}
        claim_audit = audit_claim_sources(outline, selection["sources"])
        if claim_audit["blocking_slides"] and not options.get("confirm_unsupported_claims"):
            raise ValueError("提纲存在不支持的主张，请修改来源或明确确认后再渲染")
        deck.status = "rendering"; db.commit()
        update_progress(record, .28, "deck_render", "正在生成可编辑 PPTX")
        filename = f"paper-report-{book.id}-{deck.id}.pptx"; path = settings.presentations_dir / filename
        render_pptx(path, book, profile, outline, selection["sources"], deck.paper_type or "discovery", render_options)
        structural = audit_pptx(path, outline)
        update_progress(record, .66, "deck_render", "Microsoft PowerPoint 正在真实渲染预览")
        visual = render_powerpoint_preview(path, deck.id)
        issues = [*structural.get("issues", []), *rights["issues"], *visual.get("issues", [])]
        qa = {"ok": structural["ok"] and claim_audit["ok"] and visual.get("ok", False),
              "issues": issues, "structural": structural, "claim_source": claim_audit,
              "coverage": selection.get("coverage", {}), "rights": rights, "visual": visual}
        manifest = {"version": 2, "created_at": datetime.now().isoformat(), "book_id": book.id,
                    "paper_type": deck.paper_type, "selection": {key: selection.get(key) for key in
                        ("chapter_ids", "chunk_ids", "resource_ids", "chapter_titles")},
                    "source_ids": [source["source_id"] for source in selection["sources"]],
                    "coverage": selection.get("coverage", {}), "rights": rights,
                    "generator": "nature-paper2ppt-adapted", "language": "zh-CN", "editable": True}
        deck.file_path = filename; deck.qa_json = json.dumps(qa, ensure_ascii=False)
        deck.manifest_json = json.dumps(manifest, ensure_ascii=False); deck.status = "done"; db.commit()
        update_progress(record, 1, "deck_render", "PPTX 与真实渲染审计已完成")
        return {"deck_id": deck.id, "slide_count": len(outline), "paper_type": deck.paper_type, "qa": qa}
    except Exception as exc:
        db.rollback(); deck = db.get(PresentationDeck, deck_id)
        if deck: deck.status = "outline_ready" if deck.outline_json else "failed"; deck.error_msg = str(exc); db.commit()
        raise
    finally:
        db.close()


async def generate_deck_task(record, deck_id: int) -> dict:
    await generate_outline_task(record, deck_id)
    db = SessionLocal()
    try:
        deck = db.get(PresentationDeck, deck_id)
        options = json.loads(deck.options_json or "{}")
        options["confirm_unsupported_claims"] = True
        deck.options_json = json.dumps(options, ensure_ascii=False); db.commit()
    finally:
        db.close()
    return await render_deck_task(record, deck_id)
