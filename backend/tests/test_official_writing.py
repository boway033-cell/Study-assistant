"""Three real skill adapters; model traffic is always mocked in tests."""
import asyncio
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.api import official_writing as api
from backend.app.api.writing import router as writing_router
from backend.app.core.database import Base, get_db
from backend.app.services import official_writing as service

BRIEF = {"title": "关于开展资料核验的通知", "genre": "通知", "facts": "示例单位拟开展资料核验。共12份材料。日期待定。",
         "issuer": "示例单位", "recipient": "各科室", "relationship": "下行", "length": 800}


@pytest.fixture
def client(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)

    def db():
        with factory() as session:
            yield session

    async def model(db, stage, payload):
        if stage == "outline":
            return "一、核验安排\n二、材料要求\n【待补：日期】"
        if stage == "review":
            return "P0：日期待补，请核对实际安排。"
        if stage == "revise":
            return "各科室：\n核验12份材料。请保留核验记录。\n【待补：日期】"
        return "各科室：\n一、核验安排\n请核验12份材料。\n【待补：日期】\n[落款]示例单位"

    monkeypatch.setattr(service, "model_text", model)
    app = FastAPI()
    app.include_router(api.router)
    app.include_router(writing_router)
    app.dependency_overrides[get_db] = db
    with TestClient(app) as c:
        yield c
    engine.dispose()


def post(client, row, action, **payload):
    return client.post(f'/api/writing/official/outputs/{row["id"]}/{action}', json={"etag": row["etag"], **payload})


def outline(client):
    result = client.post("/api/writing/official/outline", json={"brief": BRIEF, "allow_model": True})
    assert result.status_code == 200, result.text
    return result.json()


def draft(client):
    row = outline(client)
    row = post(client, row, "confirm", title=row["title"], text=row["text"]).json()
    result = post(client, row, "draft", allow_model=True)
    assert result.status_code == 200, result.text
    return result.json()


def test_confirmation_and_consent_are_server_gates(client):
    assert client.post("/api/writing/official/outline", json={"brief": BRIEF, "allow_model": False}).status_code == 422
    row = outline(client)
    assert post(client, row, "draft", allow_model=True).status_code == 409
    confirmed = post(client, row, "confirm", title=row["title"], text=row["text"]).json()
    assert confirmed["id"] != row["id"] and confirmed["audit"]["confirmed"]
    saved = post(client, confirmed, "save", title=row["title"], text="修改后的提纲").json()
    assert not saved["audit"]["confirmed"]
    assert post(client, saved, "draft", allow_model=True).status_code == 409
    assert post(client, confirmed, "draft", allow_model=False).status_code == 422
    assert post(client, confirmed, "export", review_acknowledged=True).status_code == 409


def test_version_integrity_review_and_legacy_protection(client):
    row = draft(client)
    reviewed = post(client, row, "review", allow_model=True).json()
    assert reviewed["text"] == row["text"] and "P0" in reviewed["audit"]["review"]
    revised = post(client, reviewed, "revise", allow_model=True, instructions="保留核验记录").json()
    assert revised["id"] != row["id"]
    assert revised["audit"]["parent_id"] == row["id"]
    assert client.get(f'/api/writing/official/outputs/{row["id"]}').json()["text"] == row["text"]
    assert client.patch(f'/api/writing/outputs/{row["id"]}', json={"title": "覆盖", "output_text": "覆盖"}).status_code == 409
    assert client.get('/api/writing/outputs?exclude_official=true').json()["total"] == 0
    assert client.get('/api/writing/official/outputs').json()["total"] == 4
    assert post(client, {**row, "etag": "0" * 64}, "save", title="改稿", text="正文").status_code == 409


def test_actual_lieflat_checker_not_invented_baselines():
    checks = service.inspect_text("请予批准。新增95%目标。【待补：日期】", {**BRIEF, "genre": "报告"})
    assert not checks["lieflat"]["applicable"]
    assert checks["placeholders"] == ["【待补：日期】"]
    assert any("请批" in warning for warning in checks["warnings"])
    measured = service.inspect_text("一、情况\n核验12份材料。\n二、问题\n部分记录待核。", {**BRIEF, "genre": "调研报告"})
    assert measured["lieflat"]["applicable"]
    assert measured["lieflat"]["sample_size"] == service.check_params.REF["调研报告"]["n"]
    metric = next(m for m in measured["lieflat"]["metrics"] if m["key"] == "chars")
    assert metric["value"] > 0
    assert not service.inspect_text("示例〔2026〕1号", BRIEF)["placeholders"]


@pytest.mark.parametrize("overrides", [
    {"page": {"size": "Letter"}}, {"issuer": "禁止保存单位"},
    {"page": {"margins_cm": {"left": 10, "right": 10}}},
    {"styles": {"body": {"size_pt": 60, "line_spacing_pt": 10}}},
    {"styles": {"body": {"first_line_chars": 20, "size_pt": 22}}},
    {"styles": {"body": {"font_cn": "Bad\nFont"}}},
    {"styles": {"body": {"space_before_pt": float("nan")}}},
])
def test_invalid_format_is_rejected(overrides):
    with pytest.raises(ValueError):
        service.checked_profile(overrides)


def test_profile_confirmation_and_conflict(client):
    initial = client.get('/api/writing/official/options').json()["format"]
    payload = {"overrides": {"global": {"bold": False}}, "etag": initial["etag"], "confirmed": False}
    assert client.put('/api/writing/official/profile', json=payload).status_code == 422
    payload["confirmed"] = True
    result = client.put('/api/writing/official/profile', json=payload)
    assert result.status_code == 200
    assert result.json()["profile"]["global"]["bold"] is False
    assert client.put('/api/writing/official/profile', json=payload).status_code == 409
    assert result.json()["overrides"] == {"global": {"bold": False}}


def test_real_sanmu_export_and_metadata(client):
    row = draft(client)
    assert post(client, row, "export", review_acknowledged=False).status_code == 422
    result = post(client, row, "export", overrides={"global": {"bold": False}}, review_acknowledged=True)
    assert result.status_code == 200, result.text
    exported = result.json()
    audit = exported["audit"]["export"]
    assert audit["validation"] == "structural_pass" and audit["visual_status"] == "not_checked"
    assert audit["classifications"]
    response = client.get(f'/api/writing/outputs/{row["id"]}/download')
    assert response.status_code == 200
    import io
    document = Document(io.BytesIO(response.content))
    assert any("12份材料" in p.text for p in document.paragraphs)
    assert document.sections[0].page_width.cm == pytest.approx(21, abs=0.01)
    with ZipFile(io.BytesIO(response.content)) as archive:
        core = archive.read("docProps/core.xml").decode()
        from lxml import etree
        root = etree.fromstring(core.encode())
        assert not root.xpath('//*[local-name()="creator" or local-name()="lastModifiedBy"]/text()')
    assert exported["text"] == row["text"]


def test_model_failure_does_not_create_output_or_leak_secrets(client, monkeypatch):
    async def fail(*args):
        raise RuntimeError("Authorization: secret-api-key")
    monkeypatch.setattr(service, "model_text", fail)
    result = client.post('/api/writing/official/outline', json={"brief": BRIEF, "allow_model": True})
    assert result.status_code == 502 and "secret-api-key" not in result.text
    assert client.get('/api/writing/official/outputs').json()["total"] == 0


def test_timeout_is_reported(client, monkeypatch):
    async def fail(*args):
        raise asyncio.TimeoutError()
    monkeypatch.setattr(service, "model_text", fail)
    assert client.post('/api/writing/official/outline', json={"brief": BRIEF, "allow_model": True}).status_code == 504


def test_actual_prompt_loads_official_skill_and_respects_router(monkeypatch):
    captured = {}
    class Provider:
        async def stream_chat(self, messages):
            captured["messages"] = messages
            yield "一、已知情况"
    monkeypatch.setattr(service, "load_llm_config", lambda db, task: {"selected": task})
    def get(name, cfg):
        captured["config"] = cfg
        return Provider()
    monkeypatch.setattr(service.LLMRouter, "get", get)
    assert asyncio.run(service.model_text(None, "outline", {"brief": BRIEF})) == "一、已知情况"
    assert service.drafting_rules() in captured["messages"][0]["content"]
    assert "不得编造" in captured["messages"][0]["content"]
    assert captured["config"] == {"selected": "writing"}
