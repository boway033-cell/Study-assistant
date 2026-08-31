from pathlib import Path

import pytest


def test_paper_type_routing():
    from backend.app.services.presentation_deck import classify_paper_type

    assert classify_paper_type("A new transformer method", "algorithm benchmark baseline architecture") == "methods"
    assert classify_paper_type("Randomized trial", "patient cohort clinical survival intervention") == "clinical"
    assert classify_paper_type("A perspective", "review perspective future directions") == "review"


def test_private_network_urls_are_rejected(monkeypatch):
    from backend.app.services.literature_access import validate_public_https_url

    monkeypatch.setattr("socket.getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("127.0.0.1", 443))])
    with pytest.raises(ValueError, match="私有|本地|保留"):
        validate_public_https_url("https://example.test/file.pdf")


def test_proxy_fake_ip_allows_domain_but_not_literal(monkeypatch):
    from backend.app.services.literature_access import validate_public_https_url

    monkeypatch.setattr("socket.getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("198.18.0.74", 443))])
    assert validate_public_https_url("https://journal.example/article") == "https://journal.example/article"
    with pytest.raises(ValueError, match="私有|本地|保留"):
        validate_public_https_url("https://198.18.0.74/article")


def test_library_handoff_does_not_embed_credentials(monkeypatch):
    from backend.app.services.literature_access import build_library_handoff

    monkeypatch.setattr("socket.getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("1.1.1.1", 443))])
    url = build_library_handoff("https://library.example/search", "10.1000/test")
    assert "q=10.1000%2Ftest" in url
    assert "cookie" not in url.lower()


def test_web_page_discovers_citation_pdf_and_relative_download(monkeypatch):
    from backend.app.services.literature_access import discover_pdf_links

    monkeypatch.setattr("socket.getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("1.1.1.1", 443))])
    source = '''<meta name="citation_pdf_url" content="/paper/full.pdf">
                <a href="/download?id=7">download</a>'''
    links = discover_pdf_links("https://journal.example/article/7", source)
    assert "https://journal.example/paper/full.pdf" in links
    assert "https://journal.example/download?id=7" in links


def test_rdfybk_page_discovers_precise_pdf_handoff(monkeypatch):
    from backend.app.services.literature_access import discover_pdf_links

    monkeypatch.setattr("socket.getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("1.1.1.1", 443))])
    links = discover_pdf_links("https://www.rdfybk.com/qw/detail?id=926513", "<script>Set_ArtIntro()</script>")
    assert links == ["https://www.rdfybk.com/qw/DownPdf?id=926513"]


def test_rdfybk_keeps_precise_handoff_when_source_blocks_probe(monkeypatch):
    import asyncio
    from backend.app.services import literature_access

    monkeypatch.setattr("socket.getaddrinfo", lambda *args, **kwargs: [(2, 1, 6, "", ("1.1.1.1", 443))])
    async def blocked(*args, **kwargs):
        raise OSError("blocked")
    monkeypatch.setattr(literature_access, "_read_probe", blocked)
    candidates = asyncio.run(literature_access._resolve_web_url("https://www.rdfybk.com/qw/detail?id=926513"))
    assert candidates[0].url == "https://www.rdfybk.com/qw/DownPdf?id=926513"
    assert candidates[0].direct_download is False


def test_default_text_router_remains_deepseek():
    from backend.app.services.llm import DeepSeekProvider, LLMRouter

    assert isinstance(LLMRouter.get().providers[0], DeepSeekProvider)


def test_presentation_freshness_detects_changed_selected_knowledge():
    from uuid import uuid4
    from backend.app.api.presentations import _source_freshness
    from backend.app.core.database import SessionLocal
    from backend.app.models import Book, KnowledgeNote
    from backend.app.services.presentation_deck import source_fingerprint

    db = SessionLocal()
    book = Book(title=f"deck-freshness-{uuid4().hex}", file_path="deck.pdf", file_type="pdf", status="ready")
    db.add(book); db.flush()
    note = KnowledgeNote(book_id=book.id, title="PPT 来源", content="原始证据", source_refs_json="[]")
    db.add(note); db.commit(); db.refresh(note)
    selection = {"sources": [{"source_id": f"note:{note.id}", "chapter_title": note.title,
                               "snapshot_hash": source_fingerprint(note.title, note.content, note.source_refs_json)}]}
    try:
        assert _source_freshness(db, selection)["status"] == "fresh"
        note.content = "更新证据"; db.commit()
        assert _source_freshness(db, selection)["status"] == "stale"
    finally:
        db.query(KnowledgeNote).filter(KnowledgeNote.id == note.id).delete(synchronize_session=False)
        db.commit()
        db.query(Book).filter(Book.id == book.id).delete(synchronize_session=False)
        db.commit(); db.close()


def test_local_outline_and_editable_pptx(tmp_path):
    from pptx import Presentation
    from backend.app.models import Book
    from backend.app.services.presentation_deck import _local_outline, audit_pptx, render_pptx

    sources = [{"source_id": "chunk:1", "chunk_id": 1, "chapter_id": 1,
                "page_start": 2, "page_end": 2, "text": "该方法提升了基准性能。边界条件仍需进一步验证。"}]
    outline = _local_outline("测试论文", "methods", sources, 8)
    book = Book(id=99, title="测试论文", file_path="missing.pdf", file_type="pdf", status="ready")
    path = tmp_path / "report.pptx"
    render_pptx(path, book, None, outline, sources, "methods",
                {"include_figures": False, "audience": "课题组", "purpose": "汇报", "slide_count": 8, "duration_minutes": 10})
    prs = Presentation(path)
    assert len(prs.slides) == 8
    assert any(shape.has_text_frame for shape in prs.slides[1].shapes)
    assert audit_pptx(path, outline)["ok"] is True


def test_claim_source_audit_flags_number_not_in_source():
    from backend.app.services.presentation_deck import audit_claim_sources

    sources = [{"source_id": "chunk:1", "chapter_id": 1, "page_start": 2,
                "text": "该方法在测试集上提升了 12%。"}]
    outline = [{"title": "封面", "kind": "cover", "claim": "", "bullets": [], "source_ids": []},
               {"title": "结果", "kind": "evidence", "claim": "性能提升 30%",
                "bullets": [], "source_ids": ["chunk:1"]}]
    audit = audit_claim_sources(outline, sources)
    assert audit["ok"] is False
    assert audit["blocking_slides"] == [2]
    assert "30%" in audit["items"][0]["missing_numbers"]


def test_stratified_selection_covers_each_selected_chapter():
    from backend.app.api.books import delete_book
    from backend.app.core.database import SessionLocal
    from backend.app.models import Book, Chapter, Chunk
    from backend.app.services.presentation_deck import collect_selection

    db = SessionLocal()
    try:
        book = Book(title="Coverage", file_path="missing.pdf", file_type="pdf", status="ready")
        db.add(book); db.flush()
        chapters = [Chapter(book_id=book.id, title=f"Chapter {i}", level=1, order_index=i) for i in range(2)]
        db.add_all(chapters); db.flush()
        for chapter in chapters:
            db.add_all([Chunk(book_id=book.id, chapter_id=chapter.id, chunk_index=chapter.order_index * 10 + i,
                              content=(f"章节{chapter.order_index}片段{i}。" * 120), page_start=i + 1, page_end=i + 1)
                        for i in range(6)])
        db.commit()
        selection = collect_selection(db, book.id, [chapter.id for chapter in chapters], [], "", [], 8000)
        assert selection["coverage"]["total_groups"] == 2
        assert selection["coverage"]["covered_groups"] == 2
        assert {source["chapter_id"] for source in selection["sources"]} == {chapter.id for chapter in chapters}
        delete_book(book.id, db)
    finally:
        db.close()


def test_pptx_requires_and_propagates_selected_knowledge_objects():
    from backend.app.api.books import delete_book
    from backend.app.core.database import SessionLocal
    from backend.app.models import Book, Chapter, Chunk, KnowledgeNote
    from backend.app.services.presentation_deck import collect_multi_selection, validate_outline

    db = SessionLocal()
    try:
        book = Book(title="Knowledge-bound deck", file_path="missing.pdf", file_type="pdf", status="ready")
        db.add(book); db.flush()
        chapter = Chapter(book_id=book.id, title="证据章", level=1, order_index=0)
        db.add(chapter); db.flush()
        db.add(Chunk(book_id=book.id, chapter_id=chapter.id, chunk_index=0,
                     content="原始文献定位材料。" * 200, page_start=1, page_end=1))
        note = KnowledgeNote(book_id=book.id, title="已审查共识", content="两项独立研究在限定条件下方向一致。")
        db.add(note); db.commit(); db.refresh(note)
        with pytest.raises(ValueError, match="至少选择一个知识对象"):
            collect_multi_selection(db, book.id, [book.id], [], [], "")
        selection = collect_multi_selection(db, book.id, [book.id], [], [], "",
                                            knowledge_note_ids=[note.id])
        assert selection["knowledge_source_ids"] == [f"note:{note.id}"]
        assert selection["sources"][0]["knowledge_type"] == "note"
        slides = [{"title": "封面", "kind": "cover", "claim": "", "bullets": [], "source_ids": []},
                  {"title": "结论", "kind": "content", "claim": "方向一致", "bullets": [],
                   "source_ids": [f"note:{note.id}"]}]
        assert validate_outline(slides, selection["sources"], selection["knowledge_source_ids"])[1]["source_ids"]
        raw_id = next(item["source_id"] for item in selection["sources"] if item.get("knowledge_type") is None)
        slides[1]["source_ids"] = [raw_id]
        with pytest.raises(ValueError, match="必须引用至少一个已选知识对象"):
            validate_outline(slides, selection["sources"], selection["knowledge_source_ids"])
        delete_book(book.id, db)
    finally:
        db.close()


def test_bookshelf_is_virtual_and_multi_membership():
    from backend.app.api.shelves import ShelfBooksWrite, ShelfWrite, create_shelf, delete_shelf, put_shelf_books
    from backend.app.api.books import delete_book
    from backend.app.core.database import SessionLocal
    from backend.app.models import Book, Shelf, shelf_books
    from sqlalchemy import select

    db = SessionLocal()
    try:
        book = Book(title="Shelf item", file_path="same.pdf", file_type="pdf", status="ready")
        db.add(book); db.commit(); db.refresh(book)
        first = create_shelf(ShelfWrite(name="项目甲"), db)
        second = create_shelf(ShelfWrite(name="课程乙"), db)
        put_shelf_books(first["id"], ShelfBooksWrite(book_ids=[book.id]), db)
        put_shelf_books(second["id"], ShelfBooksWrite(book_ids=[book.id]), db)
        assert db.scalars(select(shelf_books.c.shelf_id).where(shelf_books.c.book_id == book.id)).all() == [first["id"], second["id"]]
        delete_shelf(first["id"], db)
        assert db.get(Book, book.id) is not None
        delete_shelf(second["id"], db); delete_book(book.id, db)
    finally:
        db.close()


def test_fallback_app_ports_can_create_shelves():
    """启动器切换到备用端口后，本地写操作仍应通过 Origin 防护。"""
    from backend.app.main import _ALLOWED_ORIGINS, health

    assert "http://127.0.0.1:8000" in _ALLOWED_ORIGINS
    assert "http://127.0.0.1:8010" in _ALLOWED_ORIGINS
    assert "http://127.0.0.1:8011" in _ALLOWED_ORIGINS
    assert "http://localhost:8010" in _ALLOWED_ORIGINS
    assert all("0.0.0.0" not in origin for origin in _ALLOWED_ORIGINS)
    payload = health()
    assert payload["app"] == "study-assistant"
    assert payload["api_revision"] >= 4
    assert payload["capabilities"]["shelves_write"] is True
    assert payload["capabilities"]["knowledge_insights"] is True


def test_runtime_launcher_rejects_backends_without_knowledge_insights():
    root = Path(__file__).resolve().parents[2]
    launcher = (root / "scripts" / "runtime" / "auto_start.ps1").read_text(encoding="utf-8")

    assert "[int]$response.api_revision -ge 4" in launcher
    assert "$response.capabilities.knowledge_insights -eq $true" in launcher


def test_public_deck_excludes_unknown_rights_figures():
    from backend.app.api.books import delete_book
    from backend.app.core.database import SessionLocal
    from backend.app.models import Book, LiteratureResource
    from backend.app.services.presentation_deck import _rights_policy

    db = SessionLocal()
    try:
        book = Book(title="Rights", file_path="rights.pdf", file_type="pdf", status="ready")
        db.add(book); db.flush()
        db.add(LiteratureResource(book_id=book.id, title="Main", role="main",
                                  rights_status="not_evaluated", allow_reuse=0))
        db.commit()
        policy = _rights_policy(db, book.id, {"use_scope": "public", "include_figures": True,
                                             "rights_acknowledged": True})
        assert policy["allow_figures"] is False
        assert policy["issues"]
        delete_book(book.id, db)
    finally:
        db.close()


def test_real_powerpoint_render_when_available(tmp_path, monkeypatch):
    import hashlib
    from backend.app.models import Book
    from backend.app.services import office_render
    from backend.app.services.powerpoint_render import powerpoint_available, render_powerpoint_preview
    from backend.app.services.presentation_deck import _local_outline, render_pptx

    if not powerpoint_available():
        pytest.skip("Microsoft PowerPoint unavailable")
    sources = [{"source_id": "chunk:1", "chunk_id": 1, "chapter_id": 1,
                "page_start": 2, "page_end": 2, "text": "真实渲染测试来源。"}]
    outline = _local_outline("真实渲染", "methods", sources, 6)
    book = Book(id=991, title="真实渲染", file_path="missing.pdf", file_type="pdf", status="ready")
    path = tmp_path / "powerpoint.pptx"
    render_pptx(path, book, None, outline, sources, "methods",
                {"include_figures": False, "audience": "课题组", "purpose": "测试", "slide_count": 6, "duration_minutes": 10})
    result = render_powerpoint_preview(path, 991, timeout=60)
    assert result["ok"] is True, result
    assert result["engine"] == "microsoft-powerpoint"
    assert result["slide_count"] == 6
    monkeypatch.setattr(office_render.settings, "structured_dir", tmp_path / "structured")
    rendered_pdf = office_render.render_office_pdf(
        path, "pptx", hashlib.sha256(path.read_bytes()).hexdigest(), timeout=60,
    )
    assert rendered_pdf.read_bytes()[:5] == b"%PDF-"


def test_powerpoint_unavailable_is_explicitly_skipped(tmp_path, monkeypatch):
    from backend.app.services import powerpoint_render

    path = tmp_path / "deck.pptx"
    path.write_bytes(b"placeholder")
    monkeypatch.setattr(powerpoint_render, "powerpoint_capability", lambda: {
        "available": False, "engine": "unavailable", "reason": "COM 不可用",
    })
    result = powerpoint_render.render_powerpoint_preview(path, 42)
    assert result["ok"] is False
    assert result["skipped"] is True
    assert result["validation_status"] == "not_rendered"


def test_complex_pdf_page_uses_fidelity_fallback(tmp_path):
    import fitz
    from backend.app.services.presentation_deck import _extract_page_visual

    source = tmp_path / "source.pdf"
    doc = fitz.open()
    page = doc.new_page()
    for index in range(14):
        page.draw_rect(fitz.Rect(20 + index, 30 + index, 180 + index, 110 + index))
    doc.save(source)
    doc.close()
    result = _extract_page_visual(source, 1, tmp_path / "visual.png", "公式 E = mc²")
    assert result is not None
    assert result["mode"] == "page_fidelity"
    assert result["formula_hint"] is True
    assert result["path"].exists()
