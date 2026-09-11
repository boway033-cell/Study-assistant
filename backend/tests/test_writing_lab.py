"""写作实验室的硬边界、信息守恒与 Word 格式测试。"""
import asyncio
import inspect
import json
import re
from pathlib import Path

import docx
import pytest
from sqlalchemy import select

from ai_tone_cases import DETECTOR_CASES, RULE_CASES
from backend.app.api.writing import (OutputReviewReq, OutputUpdateReq, list_outputs, review_output,
                                     update_output)
from backend.app.core.config import settings
from backend.app.core.database import SessionLocal
from backend.app.models import WritingOutput
from backend.app.services import writing_lab
from backend.app.services.writing_lab import (
    AI_TONE_VIOLATION_PATTERNS,
    _replace_across_runs,
    _validate_replacement,
    ai_flavor_violations,
    apply_text_changes,
    clean_blocks,
    create_word_output,
    text_blocks,
    validate_corpus,
    collect_writing_knowledge,
    collect_literature_evidence,
    build_literature_review_prompt,
)
from backend.app.services.writing_citations import readable_citations


def test_writing_output_archive_is_single_entry_point_with_kind_filter():
    """写作输出是唯一归档端：默认排除公文，且可按类型筛选。"""
    db = SessionLocal()
    try:
        everything = list_outputs(1, 100, True, None, db)
        assert all(row["kind"] != "official_document" for row in everything["items"])
        assert set(row["kind"] for row in everything["items"]).issubset(
            {"imitation", "literature_review", "ai_tone"})
        tone_only = list_outputs(1, 100, True, "ai_tone", db)
        assert all(row["kind"] == "ai_tone" for row in tone_only["items"])
        assert tone_only["total"] <= everything["total"]
    finally:
        db.close()


def test_writing_dna_rejects_fewer_than_twenty_articles():
    with pytest.raises(ValueError, match="至少需要 20 篇"):
        validate_corpus(None, list(range(1, 20)))


def test_writing_material_requires_selected_in_scope_knowledge_objects():
    from uuid import uuid4
    from backend.app.models import Book, EvidenceCard, KnowledgeNote

    db = SessionLocal()
    book = Book(title=f"writing-source-{uuid4().hex}", file_path="writing.pdf", file_type="pdf", status="ready")
    db.add(book); db.commit(); db.refresh(book)
    note = KnowledgeNote(book_id=book.id, title="已选笔记", content="只允许使用这条事实", origin="user")
    card = EvidenceCard(book_id=book.id, title="冲突证据", evidence_text="结果方向相反",
                        claim_text="不同研究存在冲突", verification_status="needs_review")
    db.add_all([note, card]); db.commit(); db.refresh(note); db.refresh(card)
    try:
        with pytest.raises(ValueError, match="至少选择一个知识对象"):
            collect_writing_knowledge(db, [book.id], [], [], [])
        context, manifest = collect_writing_knowledge(db, [book.id], [note.id], [card.id], [])
        assert f"[NOTE:{note.id}|B{book.id}]" in context
        assert f"[EVIDENCE:{card.id}|B{book.id}]" in context
        assert {item["type"] for item in manifest} == {"note", "evidence"}
    finally:
        db.delete(book); db.commit(); db.close()


def test_writing_output_freshness_detects_changed_and_missing_knowledge():
    from uuid import uuid4
    from backend.app.api.writing import _source_freshness
    from backend.app.models import Book, KnowledgeNote

    db = SessionLocal()
    book = Book(title=f"freshness-{uuid4().hex}", file_path="fresh.pdf", file_type="pdf", status="ready")
    db.add(book); db.flush()
    note = KnowledgeNote(book_id=book.id, title="来源", content="原始判断", source_refs_json="[]")
    db.add(note); db.commit(); db.refresh(note)
    _, manifest = collect_writing_knowledge(db, [book.id], [note.id], [], [])
    output = WritingOutput(kind="imitation", title="快照", input_type="text", output_text="正文",
                           audit_json=json.dumps({"knowledge_objects": manifest}, ensure_ascii=False))
    db.add(output); db.commit(); db.refresh(output)
    try:
        assert _source_freshness(db, output)["status"] == "fresh"
        note.content = "修订后的判断"; db.commit()
        assert _source_freshness(db, output)["status"] == "stale"
        db.delete(note); db.commit()
        assert _source_freshness(db, output)["status"] == "missing"
    finally:
        db.delete(output); db.delete(book); db.commit(); db.close()


def test_literature_review_uses_fair_selected_document_packets_with_stable_anchors():
    from uuid import uuid4
    from backend.app.models import Book, Chunk

    db = SessionLocal()
    books = []
    try:
        for book_index, label in enumerate(("社科", "文学", "生物"), start=1):
            book = Book(title=f"{label}-{uuid4().hex}", file_path=f"review-{book_index}.pdf",
                        file_type="pdf", status="ready", total_pages=12)
            db.add(book); db.flush(); books.append(book)
            for chunk_index in range(10):
                db.add(Chunk(book_id=book.id, chunk_index=chunk_index,
                             page_start=chunk_index + 1, page_end=chunk_index + 1,
                             content=f"{label}材料第{chunk_index}段。共同议题与不同证据。" * 30))
        db.commit()
        context, manifest, anchors = collect_literature_evidence(
            db, [book.id for book in books], "共同议题中的冲突与互补证据", max_chars=15000,
        )
        assert len(manifest) == 3
        assert all(item["selected_chunks"] <= 8 for item in manifest)
        assert all(any(anchor.startswith(f"[B{book.id}:") for anchor in anchors) for book in books)
        assert all(f"B{book.id}《{book.title}》" in context for book in books)
    finally:
        for book in books:
            db.delete(book)
        db.commit(); db.close()


@pytest.mark.parametrize("discipline, expected", [
    ("social_science", "因果识别与相关性边界"),
    ("humanities", "竞争性阐释和反向阅读"),
    ("natural_biomedical", "观察性关联不得改写为因果"),
])
def test_literature_review_prompt_has_domain_lens_dna_tone_and_method_boundary(discipline, expected):
    prompt = build_literature_review_prompt(
        "证据为何冲突？", "跨文献综述", "narrative", discipline, 3000,
        "[B1:C1:P1-2]\n证据一\n\n[B2:C2:P3-3]\n证据二",
        "短句与克制表达", True,
    )
    assert expected in prompt
    assert "按主题综合" in prompt
    assert "自由建立概念联系" in prompt
    assert "有中心判断、论证连续" in prompt
    assert "系统会在成文后自动转成脚注编号" in prompt
    assert "免责声明和修饰词堆叠" in prompt
    assert "Writing DNA：只约束表达" in prompt
    assert "去 AI 味生成约束" in prompt
    assert "禁止自称“系统综述”“元分析”或“PRISMA 合规”" in prompt


@pytest.mark.asyncio
async def test_literature_review_generation_persists_citation_and_scope_audit(monkeypatch):
    from uuid import uuid4
    from backend.app.models import Book, Chunk
    from backend.app.services import writing_lab

    db = SessionLocal()
    books = []
    output = None
    try:
        chunks = []
        for index in range(2):
            book = Book(title=f"review-generation-{uuid4().hex}", file_path=f"generation-{index}.pdf",
                        file_type="pdf", status="ready", total_pages=3)
            db.add(book); db.flush(); books.append(book)
            chunk = Chunk(book_id=book.id, chunk_index=0, page_start=1, page_end=2,
                          content="这是可核验的研究材料，保留冲突、限定条件与证据边界。" * 40)
            db.add(chunk); db.flush(); chunks.append(chunk)
        db.commit()

        class FakeProvider:
            async def stream_chat(self, messages):
                assert "Writing DNA：只约束表达" in messages[-1]["content"]
                yield (f"# 综述\n\n两篇材料形成互补证据 {_anchor(chunks[0])} {_anchor(chunks[1])}\n\n"
                       "## 结论\n仍需人工核对来源边界。")

        def _anchor(chunk):
            return f"[B{chunk.book_id}:C{chunk.id}:P1-2]"

        monkeypatch.setattr(writing_lab, "load_llm_config", lambda db, task: {})
        monkeypatch.setattr(writing_lab.LLMRouter, "get", staticmethod(lambda mode, cfg: FakeProvider()))
        output = await writing_lab.generate_literature_review(
            db, question="两篇材料如何互补？", title="测试综述",
            book_ids=[book.id for book in books], discipline="auto",
        )
        audit = json.loads(output.audit_json)
        assert output.kind == "literature_review"
        assert audit["content_policy"] == "selected_library_documents_only"
        assert audit["cited_book_ids"] == [book.id for book in books]
        assert audit["unreferenced_book_ids"] == []
        assert "[B" not in output.output_text
        assert "## 来源索引" in output.output_text
        assert audit["citation_notes"][0]["anchor"].startswith("B")
    finally:
        if output is not None:
            db.delete(output)
        for book in books:
            db.delete(book)
        db.commit(); db.close()


@pytest.mark.asyncio
async def test_literature_review_keeps_draft_when_citation_coverage_needs_review(monkeypatch):
    from uuid import uuid4
    from backend.app.models import Book, Chunk
    from backend.app.services import writing_lab

    db = SessionLocal(); books = []; output = None
    try:
        for index in range(2):
            book = Book(title=f"review-warning-{uuid4().hex}", file_path=f"warning-{index}.pdf",
                        file_type="pdf", status="ready", total_pages=2)
            db.add(book); db.flush(); books.append(book)
            db.add(Chunk(book_id=book.id, chunk_index=0, page_start=1, page_end=1,
                         content="用于跨文献综合的有效材料。" * 80))
        db.commit()
        first = db.scalars(select(Chunk).where(Chunk.book_id == books[0].id)).first()

        class FakeProvider:
            async def stream_chat(self, messages):
                yield f"# 草稿\n\n保留可用结果 {_anchor(first)} [B999:C999:P1-1]"

        def _anchor(chunk):
            return f"[B{chunk.book_id}:C{chunk.id}:P1-1]"

        monkeypatch.setattr(writing_lab, "load_llm_config", lambda db, task: {})
        monkeypatch.setattr(writing_lab.LLMRouter, "get", staticmethod(lambda mode, cfg: FakeProvider()))
        output = await writing_lab.generate_literature_review(
            db, question="如何综合这些材料？", title="保留草稿", book_ids=[book.id for book in books])
        audit = json.loads(output.audit_json)
        assert output.output_text.startswith("# 草稿")
        assert audit["invalid_anchors"] == ["[B999:C999:P1-1]"]
        assert audit["citation_warning"]
        assert books[1].id in audit["unreferenced_book_ids"]
        assert "[B999" not in output.output_text
    finally:
        if output is not None:
            db.delete(output)
        for book in books:
            db.delete(book)
        db.commit(); db.close()


def test_machine_anchors_become_compact_reader_notes_without_losing_audit_map():
    source = "核心判断 [B70:C12770:P12-12]。进一步解释 [B70:C12764:P9-9]。"
    rendered, notes = readable_citations(
        source,
        valid_anchors={"B70:C12770:P12-12", "B70:C12764:P9-9"},
        labels={"B70:C12770:P12-12": "《测试文献》，第 12 页",
                "B70:C12764:P9-9": "《测试文献》，第 9 页"},
    )
    assert "[B70" not in rendered
    assert "核心判断 ¹。" in rendered
    assert "## 来源索引" in rendered
    assert notes[0] == {"number": 1, "anchor": "B70:C12770:P12-12", "label": "《测试文献》，第 12 页"}


def test_text_change_preserves_line_structure():
    source = "标题\n\n说白了，结果可能是 12%。\n末段\n"
    changes = [{"id": "line:2", "old": "说白了，", "new": "", "rule_ids": [9]}]
    assert apply_text_changes(source, changes) == "标题\n\n结果可能是 12%。\n末段\n"
    assert [item["id"] for item in text_blocks(source)] == ["line:0", "line:2", "line:3"]


def test_information_guard_rejects_number_and_qualifier_loss():
    with pytest.raises(ValueError, match="12%"):
        _validate_replacement("结果可能提高 12%", "结果可能提高")
    with pytest.raises(ValueError, match="限定词"):
        _validate_replacement("结果可能提高", "结果提高")


def test_cross_run_replacement_keeps_unaffected_run_formatting():
    document = docx.Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("说白")
    first.bold = True
    second = paragraph.add_run("了，结论")
    second.italic = True
    assert _replace_across_runs(paragraph, "说白了，", "") is True
    assert paragraph.text == "结论"
    assert second.text == "结论"
    assert second.italic is True


def test_word_output_uses_stzhongsong_as_primary_font(tmp_path: Path):
    destination = tmp_path / "writing.docx"
    create_word_output("测试", "# 标题\n\n正文", destination)
    document = docx.Document(destination)
    normal_xml = document.styles["Normal"]._element.xml
    assert "华文中宋" in normal_xml
    assert document.paragraphs[0].text == "测试"
    assert "华文中宋" in document.styles["Title"]._element.xml
    assert destination.stat().st_size > 0


def test_ai_tone_review_can_restore_or_apply_each_change():
    db = SessionLocal()
    try:
        row = WritingOutput(kind="ai_tone", title="审阅稿", input_type="text",
                            source_text="说白了，结果可能提高 12%。", output_text="结果可能提高 12%。",
                            audit_json='{"changes":[{"id":"line:0","old":"说白了，","new":"","rule_ids":[9]}]}')
        db.add(row); db.commit(); db.refresh(row)
        restored = review_output(row.id, OutputReviewReq(accepted_indexes=[]), db)
        assert restored["output_text"] == "说白了，结果可能提高 12%。"
        filename = db.get(WritingOutput, row.id).output_file_path
        applied = review_output(row.id, OutputReviewReq(accepted_indexes=[0]), db)
        assert applied["output_text"] == "结果可能提高 12%。"
        assert applied["audit"]["accepted_indexes"] == [0]
        assert db.get(WritingOutput, row.id).output_file_path == filename
    finally:
        db.close()


def test_imitation_output_can_be_edited_and_regenerates_same_word_file():
    db = SessionLocal()
    try:
        row = WritingOutput(kind="imitation", title="初稿", input_type="text",
                            source_text="", output_text="原内容", audit_json="{}")
        db.add(row); db.commit(); db.refresh(row)
        first = update_output(row.id, OutputUpdateReq(title="修订稿", output_text="修订内容"), db)
        filename = db.get(WritingOutput, row.id).output_file_path
        second = update_output(row.id, OutputUpdateReq(title="再修订", output_text="最终内容"), db)
        assert first["audit"]["user_edited"] is True
        assert second["output_text"] == "最终内容"
        assert db.get(WritingOutput, row.id).output_file_path == filename
    finally:
        db.close()


def test_docx_review_replays_selected_changes_from_original_format():
    source_dir = settings.writing_dir / "sources"
    source_dir.mkdir(parents=True, exist_ok=True)
    source = source_dir / "review-format-source.docx"
    document = docx.Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("说白"); first.bold = True
    second = paragraph.add_run("了，结论"); second.italic = True
    document.save(source)
    db = SessionLocal()
    try:
        row = WritingOutput(kind="ai_tone", title="Word审阅稿", input_type="docx",
                            source_text=None, output_text="结论", source_file_path=source.name,
                            audit_json='{"changes":[{"id":"p:0","old":"说白了，","new":"","rule_ids":[9]}]}')
        db.add(row); db.commit(); db.refresh(row)
        result = review_output(row.id, OutputReviewReq(accepted_indexes=[0]), db)
        saved = docx.Document(settings.writing_dir / "outputs" / db.get(WritingOutput, row.id).output_file_path)
        assert result["audit"]["applied"] == 1
        assert saved.paragraphs[0].text == "结论"
        assert saved.paragraphs[0].runs[-1].italic is True
    finally:
        db.close()


def test_ai_flavor_postposed_attr_detects_targeted_patterns():
    """「禁止定语后置」必须命中 D9 四类真实句式与原有「，……的。」形态。"""
    hits = ai_flavor_violations(
        "治理的有效性，其根源在于制度设计。"
        "责任分配的效率，其决定因素是考核方式。"
        "这项工作的结果，关键在于执行力度。"
        "该政策的成败，取决于基层配合。"
        "结果，取决于执行的。"
    )
    assert hits.get("postposed_attr"), "应检测到定语后置违规"
    snippets = " ".join(hits["postposed_attr"])
    assert "其根源在于" in snippets, "未命中「X的Y，其根源在于Z」"
    assert "其决定因素是" in snippets, "未命中「X的Y，其决定因素是Z」"
    assert "关键在于" in snippets, "未命中「X的Y，关键在于Z」"
    assert "取决于" in snippets, "未命中「X的Y，取决于Z」"
    assert "取决于执行的" in snippets, "未保留原「，……的。」形态"


def test_ai_flavor_postposed_attr_skips_judgment_sentences():
    """「是……的」正常判断句不得被误报为定语后置。"""
    for sentence in ("其结果是稳定的。", "数据是可靠的。", "治理改革的结果是稳定的。"):
        violations = ai_flavor_violations(sentence)
        assert not violations.get("postposed_attr"), f"不应误报判断句：{sentence}"


# ---------------------------------------------------------------------------
# AI 味禁令测试网：正例 / 反例的唯一真源在 ai_tone_cases.py，
# 以下元测试保证「规则声明」与「用例」不会再各自漂移（D9 静默失效的根因）。
# ---------------------------------------------------------------------------


def _prompt_rule_ids() -> set[int]:
    """从 AI_TONE_RULES 文本里解析出实际声明的规则编号。"""
    return {int(value) for value in
            re.findall(r"(?:^|[；\n])\s*(\d{1,2})\s(?=\S)", writing_lab.AI_TONE_RULES, re.M)}


@pytest.mark.parametrize("detector", sorted(DETECTOR_CASES))
def test_ai_flavor_detector_positive_and_negative(detector):
    """四个检测器都必须做到：正例命中、反例不误报。"""
    for text in DETECTOR_CASES[detector]["positive"]:
        hits = ai_flavor_violations(text)
        assert hits.get(detector), f"{detector} 应命中：{text}"
    for text in DETECTOR_CASES[detector]["negative"]:
        hits = ai_flavor_violations(text)
        assert not hits.get(detector), f"{detector} 不应误报：{text}"


def test_ai_flavor_detector_keys_stay_in_sync():
    """fixture 的检测器 key 必须与代码里的检测器全集完全一致（双向）。"""
    expected = set(AI_TONE_VIOLATION_PATTERNS) | {"postposed_attr", "repeated_example_pairs"}
    assert set(DETECTOR_CASES) == expected, "fixture 与 ai_flavor_violations 的 key 集合不一致"


def test_whitelist_rule_cases_cover_every_declared_rule():
    """白名单声明的每个规则编号都必须有正反例，且不得出现多余编号。"""
    assert _prompt_rule_ids() == set(RULE_CASES), "AI_TONE_RULES 的规则编号与用例不一致"
    for rule_id, case in RULE_CASES.items():
        # new 允许为空串（该规则就是「删去这个成分」），其余字段不得为空。
        for field in ("text", "old", "negative"):
            assert str(case.get(field) or "") != "", f"规则 {rule_id} 缺少 {field}"
        assert case.get("new") is not None, f"规则 {rule_id} 缺少 new"
        assert case["old"] in case["text"], f"规则 {rule_id} 的 old 不是 text 的精确子串"


def test_clean_blocks_rule_bound_matches_declared_rules():
    """clean_blocks 里硬编码的规则编号上界必须等于白名单实际条数。

    上界写死在 `rule > 13` 的判定里：新增规则却忘了改这里，新规则会被整体静默丢弃。
    """
    bound = re.search(r"rule > (\d+)", inspect.getsource(clean_blocks))
    assert bound, "未找到 clean_blocks 的规则编号上界判定"
    assert int(bound.group(1)) == max(RULE_CASES), "规则编号上界与白名单条数不一致"


class _StubProvider:
    """把固定 JSON 当作模型输出，用来验证白名单改写的接线，不依赖真实 LLM。"""

    def __init__(self, payload: dict):
        self.payload = payload

    async def stream_chat(self, messages):
        yield json.dumps(self.payload, ensure_ascii=False)


def _patch_stub_llm(monkeypatch, payload: dict) -> None:
    monkeypatch.setattr(writing_lab, "load_llm_config", lambda *args, **kwargs: {})
    monkeypatch.setattr(writing_lab.LLMRouter, "get", lambda *args, **kwargs: _StubProvider(payload))


@pytest.mark.parametrize("rule_id", sorted(RULE_CASES))
def test_clean_blocks_accepts_each_whitelist_rule(rule_id, monkeypatch):
    """13 条白名单规则逐条打通：改动被接受、规则编号被回报。"""
    case = RULE_CASES[rule_id]
    _patch_stub_llm(monkeypatch, {"changes": [
        {"id": "p0", "old": case["old"], "new": case["new"], "rule_ids": [rule_id]},
    ]})
    accepted, audit = asyncio.run(clean_blocks(None, [{"id": "p0", "text": case["text"]}], None))
    assert accepted == [{"id": "p0", "old": case["old"], "new": case["new"], "rule_ids": [rule_id]}]
    assert audit["rules"] == [rule_id]
    assert audit["rejected"] == []


def test_clean_blocks_rejects_out_of_range_rule_id(monkeypatch):
    """白名单以外的规则编号必须被拒绝，不能悄悄通过。"""
    _patch_stub_llm(monkeypatch, {"changes": [
        {"id": "p0", "old": "说白了，", "new": "", "rule_ids": [99]},
    ]})
    accepted, audit = asyncio.run(
        clean_blocks(None, [{"id": "p0", "text": "说白了，结论就是样本量不足。"}], None))
    assert accepted == []
    assert audit["rejected"], "越界规则编号应进入 rejected"


def test_clean_blocks_rejects_change_that_loses_protected_information(monkeypatch):
    """改写必须保持数字与限定词不变，否则拒绝（信息守恒闸门）。"""
    _patch_stub_llm(monkeypatch, {"changes": [
        {"id": "p0", "old": "团队把检索耗时从 120ms 降到 80ms", "new": "团队显著优化了检索性能", "rule_ids": [8]},
    ]})
    accepted, audit = asyncio.run(
        clean_blocks(None, [{"id": "p0", "text": "团队把检索耗时从 120ms 降到 80ms。"}], None))
    assert accepted == []
    assert audit["rejected"], "丢失数字的改写应进入 rejected"


@pytest.mark.parametrize("path", [
    "/api/writing/literature-review",
    "/api/writing/clean-text",
    "/api/writing/clean-docx",
    "/api/writing/profiles/{profile_id}/imitate",
])
def test_long_writing_endpoints_are_background_tasks(path):
    """四个长任务端点必须都是 202 + task_id。

    同步 → 202 是破坏性契约变更（前端必须同批发布）；这条断言用来防止
    将来有人把某个端点改回同步，导致前端拿不到 task_id 而静默挂死。
    """
    from backend.app.api.writing import router
    route = next((item for item in router.routes if getattr(item, "path", "") == path), None)
    assert route is not None, f"未找到路由 {path}"
    assert route.status_code == 202, f"{path} 应返回 202"
